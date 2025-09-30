import sqlite3
import pandas as pd
import numpy as np
from datetime import datetime, timedelta
import random
import requests
import json
import re
import openpyxl
from openpyxl.chart import BarChart, PieChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment
from pathlib import Path
import argparse
import sys
import logging

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

class DatabaseManager:    
    def __init__(self, db_path: str = "bank_data.db"):
        self.db_path = db_path
    
    def get_connection(self):
        conn = sqlite3.connect(self.db_path, check_same_thread=False)
        conn.row_factory = sqlite3.Row
        return conn
    
    def create_tables(self):
        conn = self.get_connection()
        cursor = conn.cursor()
        
        cursor.execute('''
        CREATE TABLE IF NOT EXISTS clients (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            birth_date DATE NOT NULL,
            region TEXT NOT NULL,
            phone TEXT,
            email TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
        ''')
        
        cursor.execute('''
        CREATE TABLE IF NOT EXISTS accounts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            client_id INTEGER NOT NULL,
            account_number TEXT UNIQUE NOT NULL,
            balance DECIMAL(15,2) NOT NULL DEFAULT 0,
            account_type TEXT NOT NULL DEFAULT 'savings',
            open_date DATE NOT NULL,
            status TEXT DEFAULT 'active',
            FOREIGN KEY (client_id) REFERENCES clients (id)
        )
        ''')
        
        cursor.execute('''
        CREATE TABLE IF NOT EXISTS transactions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            account_id INTEGER NOT NULL,
            amount DECIMAL(15,2) NOT NULL,
            date TIMESTAMP NOT NULL,
            type TEXT NOT NULL,
            description TEXT,
            reference_number TEXT UNIQUE,
            FOREIGN KEY (account_id) REFERENCES accounts (id)
        )
        ''')
        
        cursor.execute('CREATE INDEX IF NOT EXISTS idx_transactions_date ON transactions(date)')
        cursor.execute('CREATE INDEX IF NOT EXISTS idx_transactions_account ON transactions(account_id)')
        cursor.execute('CREATE INDEX IF NOT EXISTS idx_clients_region ON clients(region)')
        cursor.execute('CREATE INDEX IF NOT EXISTS idx_accounts_client ON accounts(client_id)')
        
        conn.commit()
        conn.close()
        logger.info("Jadvallar yaratildi")
    
    def generate_mock_data(self, num_clients: int = 50000):
        """Generate mock data"""
        conn = self.get_connection()
        cursor = conn.cursor()
        
        cursor.execute("SELECT COUNT(*) FROM clients")
        if cursor.fetchone()[0] > 0:
            logger.info("Ma'lumotlar allaqachon mavjud")
            conn.close()
            return
        
        logger.info("Mock data yaratish boshlandi...")
        
        regions = ['Toshkent', 'Samarqand', 'Buxoro', 'Andijon', 'Farg\'ona', 
                  'Namangan', 'Qashqadaryo', 'Surxondaryo', 'Jizzax', 'Sirdaryo',
                  'Navoiy', 'Xorazm', 'Qoraqalpog\'iston']
        
        clients_data = []
        for i in range(num_clients):
            name = self._generate_name()
            birth_date = self._random_date(datetime(1950, 1, 1), datetime(2005, 12, 31))
            region = random.choice(regions)
            phone = f"+998{random.randint(10, 99)}{random.randint(1000000, 9999999)}"
            email = f"user{i}@email.uz"
            clients_data.append((name, birth_date, region, phone, email))
            
            if i % 10000 == 0:
                logger.info(f"Mijozlar: {i}/{num_clients}")
        
        cursor.executemany('INSERT INTO clients (name, birth_date, region, phone, email) VALUES (?, ?, ?, ?, ?)', clients_data)
        
        logger.info("Hisoblar yaratilmoqda...")
        accounts_data = []
        account_types = ['savings', 'checking', 'business', 'credit']
        
        for client_id in range(1, num_clients + 1):
            num_accounts = random.randint(1, 3)
            for _ in range(num_accounts):
                account_number = f"8600{random.randint(1000000000000000, 9999999999999999)}"
                balance = round(random.uniform(1000, 100000000), 2)
                account_type = random.choice(account_types)
                open_date = self._random_date(datetime(2020, 1, 1), datetime(2024, 12, 31))
                accounts_data.append((client_id, account_number, balance, account_type, open_date))
        
        cursor.executemany('INSERT INTO accounts (client_id, account_number, balance, account_type, open_date) VALUES (?, ?, ?, ?, ?)', accounts_data)
        
        logger.info("Tranzaksiyalar yaratilmoqda...")
        cursor.execute("SELECT id FROM accounts")
        account_ids = [row[0] for row in cursor.fetchall()]
        
        transaction_types = ['debit', 'credit', 'transfer', 'payment', 'withdrawal', 'deposit']
        transactions_batch = []
        batch_size = 10000
        total = 0
        
        for account_id in account_ids:
            num_trans = random.randint(10, 50)
            for _ in range(num_trans):
                amount = round(random.uniform(-50000, 100000), 2)
                date = self._random_date(datetime(2023, 1, 1), datetime(2024, 9, 27))
                tx_type = random.choice(transaction_types)
                description = f"{tx_type.title()} transaction"
                ref = f"TX{random.randint(100000000, 999999999)}"
                transactions_batch.append((account_id, amount, date, tx_type, description, ref))
                total += 1
                
                if len(transactions_batch) >= batch_size:
                    cursor.executemany('INSERT INTO transactions (account_id, amount, date, type, description, reference_number) VALUES (?, ?, ?, ?, ?, ?)', transactions_batch)
                    conn.commit()
                    transactions_batch = []
                    logger.info(f"Tranzaksiyalar: {total}")
        
        if transactions_batch:
            cursor.executemany('INSERT INTO transactions (account_id, amount, date, type, description, reference_number) VALUES (?, ?, ?, ?, ?, ?)', transactions_batch)
        
        conn.commit()
        conn.close()
        logger.info(f"Mock data yaratish tugadi. Jami: {total}")
    
    def _generate_name(self):
        first = random.choice(['Akbar', 'Ali', 'Bobur', 'Davron', 'Eldor', 'Farrux', 'Jasur', 'Karim', 'Laziz', 'Mansur', 'Nodir', 'Otabek', 'Rustam', 'Sanjar', 'Timur', 'Aida', 'Barno', 'Dildora', 'Elnora', 'Feruza', 'Gulnora', 'Hilola', 'Iroda', 'Jamila', 'Kamola', 'Laylo', 'Malika', 'Nargiza', 'Oysha'])
        last = random.choice(['Aliyev', 'Karimov', 'Rahimov', 'Nazarov', 'Mamatov', 'Toshev', 'Safarov', 'Jumayev', 'Ergashev', 'Mirzayev', 'Komilov', 'Yunusov'])
        return f"{first} {last}"
    
    def _random_date(self, start, end):
        delta = end - start
        return (start + timedelta(days=random.randint(0, delta.days))).strftime('%Y-%m-%d')

class LLMQueryGenerator:
    """LLM query generator"""
    
    def __init__(self, model_url: str = "http://localhost:11434"):
        self.model_url = model_url
        self.model_name = "llama3.2"
    
    def generate_sql(self, prompt: str) -> str:
        schema = """
        Tables:
        1. clients (id, name, birth_date, region, phone, email)
        2. accounts (id, client_id, account_number, balance, account_type, open_date, status)
        3. transactions (id, account_id, amount, date, type, description)
        
        Regions: Toshkent, Samarqand, Buxoro, Andijon, Farg'ona, Namangan, Qashqadaryo, Surxondaryo, Jizzax, Sirdaryo, Navoiy, Xorazm, Qoraqalpog'iston
        """
        
        full_prompt = f"{schema}\n\nUser: {prompt}\n\nGenerate SQL query (only SQL, no explanation):\n"
        
        try:
            response = requests.post(f"{self.model_url}/api/generate", json={"model": self.model_name, "prompt": full_prompt, "stream": False}, timeout=30)
            
            if response.status_code == 200:
                sql = response.json().get("response", "").strip()
                return self._clean_sql(sql)
            else:
                return self._fallback_sql(prompt)
        except:
            return self._fallback_sql(prompt)
    
    def _clean_sql(self, sql):
        sql = re.sub(r'```sql\n?', '', sql)
        sql = re.sub(r'```\n?', '', sql)
        sql = ' '.join(sql.split())
        if not sql.strip().endswith(';'):
            sql += ';'
        return sql.strip()
    
    def _fallback_sql(self, prompt):
        prompt_lower = prompt.lower()
        if 'toshkent' in prompt_lower and 'mijoz' in prompt_lower:
            return "SELECT COUNT(*) as mijozlar FROM clients WHERE region = 'Toshkent';"
        elif 'viloyat' in prompt_lower:
            return "SELECT region, COUNT(*) as mijozlar FROM clients GROUP BY region ORDER BY mijozlar DESC;"
        elif 'balans' in prompt_lower:
            return "SELECT account_number, balance FROM accounts ORDER BY balance DESC LIMIT 10;"
        else:
            return "SELECT COUNT(*) as jami FROM clients;"

class ExcelExporter:    
    def export_data(self, data, filename, chart_type='bar'):
        if not data:
            return None
        
        df = pd.DataFrame(data)
        Path("reports").mkdir(exist_ok=True)
        filepath = f"reports/{filename}"
        
        with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='Data', index=False)
            workbook = writer.book
            worksheet = writer.sheets['Data']
            
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            for cell in worksheet[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center")
            
            for column in worksheet.columns:
                max_length = 0
                column_letter = column[0].column_letter
                for cell in column:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                worksheet.column_dimensions[column_letter].width = min(max_length + 2, 50)
            
            if len(df.columns) >= 2:
                rows = len(df) + 1
                try:
                    if chart_type == 'pie':
                        chart = PieChart()
                        data_ref = Reference(worksheet, min_col=2, min_row=1, max_row=rows)
                        labels = Reference(worksheet, min_col=1, min_row=2, max_row=rows)
                        chart.add_data(data_ref, titles_from_data=True)
                        chart.set_categories(labels)
                    elif chart_type == 'line':
                        chart = LineChart()
                        data_ref = Reference(worksheet, min_col=2, min_row=1, max_row=rows)
                        labels = Reference(worksheet, min_col=1, min_row=2, max_row=rows)
                        chart.add_data(data_ref, titles_from_data=True)
                        chart.set_categories(labels)
                    else:
                        chart = BarChart()
                        data_ref = Reference(worksheet, min_col=2, min_row=1, max_row=rows)
                        labels = Reference(worksheet, min_col=1, min_row=2, max_row=rows)
                        chart.add_data(data_ref, titles_from_data=True)
                        chart.set_categories(labels)
                    
                    chart.title = "Ma'lumotlar tahlili"
                    chart.width = 15
                    chart.height = 10
                    worksheet.add_chart(chart, "E2")
                except:
                    pass
        
        logger.info(f"Excel yaratildi: {filepath}")
        return filepath

class BankAnalystAssistant:    
    def __init__(self, db_path="bank_data.db"):
        self.db_manager = DatabaseManager(db_path)
        self.llm_generator = LLMQueryGenerator()
        self.excel_exporter = ExcelExporter()
    
    def setup_database(self, generate_data=True):
        logger.info("Database sozlanmoqda...")
        self.db_manager.create_tables()
        if generate_data:
            self.db_manager.generate_mock_data()
    
    def process_query(self, prompt):
        logger.info(f"Query: {prompt}")
        sql = self.llm_generator.generate_sql(prompt)
        logger.info(f"SQL: {sql}")
        
        try:
            conn = self.db_manager.get_connection()
            cursor = conn.cursor()
            cursor.execute(sql)
            results = cursor.fetchall()
            columns = [d[0] for d in cursor.description]
            data = [dict(zip(columns, row)) for row in results]
            conn.close()
            
            return {'success': True, 'sql_query': sql, 'data': data, 'row_count': len(data)}
        except Exception as e:
            logger.error(f"SQL xato: {e}")
            return {'success': False, 'error': str(e), 'sql_query': sql, 'data': []}
    
    def generate_report(self, prompt, chart_type='bar'):
        result = self.process_query(prompt)
        if not result['success'] or not result['data']:
            return None
        
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"bank_report_{timestamp}.xlsx"
        return self.excel_exporter.export_data(result['data'], filename, chart_type)

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--setup', action='store_true')
    parser.add_argument('--query', type=str)
    parser.add_argument('--chart', choices=['bar', 'pie', 'line'], default='bar')
    args = parser.parse_args()
    
    assistant = BankAnalystAssistant()
    
    if args.setup:
        assistant.setup_database(generate_data=True)
        return
    
    if args.query:
        filepath = assistant.generate_report(args.query, args.chart)
        if filepath:
            print(f"✅ Hisobot: {filepath}")
        else:
            print("❌ Xato")

if __name__ == "__main__":
    main()